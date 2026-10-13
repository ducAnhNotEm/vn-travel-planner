import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set page margins (1 inch / 2.54 cm all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TÀI LIỆU ĐẶC TẢ TÁC NHÂN VÀ YÊU CẦU HỆ THỐNG")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Đề tài: Hệ thống Lập kế hoạch Du lịch Thông minh Đa phương thức tại Việt Nam (VNTravel AI)")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SECTION I: ACTORS
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1_run = h1.add_run("I. CÁC TÁC NHÂN CỦA HỆ THỐNG (SYSTEM ACTORS)")
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(29, 78, 216) # Blue 700

    # 1. Tác nhân con người
    h2_1 = doc.add_paragraph()
    h2_1.paragraph_format.space_before = Pt(6)
    h2_1.paragraph_format.space_after = Pt(4)
    h2_1_run = h2_1.add_run("1. Tác nhân con người (Human Actors)")
    h2_1_run.font.size = Pt(12.5)
    h2_1_run.font.bold = True

    actors_human = [
        ("Khách du lịch tự do (End User / Traveler)", 
         "Người dùng trực tiếp sử dụng ứng dụng để tìm kiếm điểm đến, nhập sở thích, tải lịch trình cá nhân hóa, tùy biến lộ trình (thêm/bớt/đổi điểm, ghim khách sạn, chia tiền nhóm)."),
        ("Trưởng nhóm du lịch (Trip Leader / Group Organizer)", 
         "Đại diện nhóm bạn hoặc gia đình lập kế hoạch cho đoàn, cấu hình số lượng thành viên, chia sẻ liên kết lịch trình, phân bổ phòng nghỉ và chia đều chi phí phương tiện/ăn uống cho cả nhóm."),
        ("Quản trị viên hệ thống (System Administrator / Data Curator)", 
         "Quản trị cơ sở dữ liệu địa điểm du lịch (travel_db.db), kiểm duyệt tọa độ, địa chỉ thực tế, hình ảnh chuẩn Wikipedia/Wikimedia, giám sát hạn ngạch API ngoại vi và theo dõi nhật ký hệ thống.")
    ]

    for title, desc in actors_human:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        r_title = p.add_run(f"{title}: ")
        r_title.bold = True
        p.add_run(desc)

    # 2. Tác nhân hệ thống ngoại vi
    h2_2 = doc.add_paragraph()
    h2_2.paragraph_format.space_before = Pt(8)
    h2_2.paragraph_format.space_after = Pt(4)
    h2_2_run = h2_2.add_run("2. Tác nhân hệ thống ngoại vi (External System Actors)")
    h2_2_run.font.size = Pt(12.5)
    h2_2_run.font.bold = True

    actors_ext = [
        ("Dịch vụ Bản đồ & Chỉ đường (Map & Routing Provider - Goong Maps API)", 
         "Cung cấp dữ liệu Geocoding, ma trận khoảng cách thời gian thực và tọa độ polyline giải mã đường đi thực tế theo mạng lưới giao thông Việt Nam."),
        ("Dịch vụ Trí tuệ nhân tạo (LLM Engine - Claude 4.6 / Gemini Flash)", 
         "Đảm nhận vai trò NLU Parser (bóc tách ý định, điểm đi/đến, sở thích) và NLG Explainer (tổng hợp lời giải thích lộ trình, tư vấn đổi khách sạn thông minh)."),
        ("Hệ thống Dữ liệu Mở & Đa phương tiện (Wikimedia Commons / Google Places CDN)", 
         "Cung cấp nguồn hình ảnh chân thực, ảnh thực địa 100% cho các danh lam thắng cảnh, di tích lịch sử và cơ sở lưu trú.")
    ]

    for title, desc in actors_ext:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        r_title = p.add_run(f"{title}: ")
        r_title.bold = True
        p.add_run(desc)

    # SECTION II: FUNCTIONAL REQUIREMENTS
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    h2_run = h2.add_run("II. 10 YÊU CẦU CHỨC NĂNG (FUNCTIONAL REQUIREMENTS - FR)")
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(29, 78, 216)

    fr_data = [
        ("FR-01", "Xử lý ngôn ngữ tự nhiên & Bóc tách ý định (NLU Parsing)",
         "Hệ thống cho phép người dùng nhập yêu cầu bằng câu tự nhiên (tiếng Việt), tự động phân tích và bóc tách chính xác các thực thể: Điểm xuất phát (origin), Điểm đến (destination), Số ngày đi, Số lượng người, Phương tiện di chuyển và Thẻ sở thích mà không bị nhầm lẫn giữa điểm đón và điểm du lịch."),
        ("FR-02", "Tự động gom cụm & Phân bổ lịch trình đa ngày (Spatial Clustering)",
         "Hệ thống tự động phân bổ các điểm tham quan, di tích, giải trí vào từng ngày dựa trên mật độ khoảng cách địa lý (sử dụng thuật toán toán học xác định), đảm bảo các điểm trong cùng một ngày nằm gần nhau để giảm thiểu tối đa thời gian di chuyển."),
        ("FR-03", "Tối ưu hóa thứ tự tham quan trong ngày (TSP Routing)",
         "Với danh sách điểm trong một ngày, hệ thống tự động giải bài toán người du lịch (TSP) để sắp xếp lộ trình tham quan ngắn nhất, bắt đầu từ điểm lưu trú (hoặc điểm xuất phát) qua các chặng tham quan và quay về nơi nghỉ ngơi."),
        ("FR-04", "Phân bổ khung giờ theo Nhịp sinh học du lịch Việt Nam (Human Rhythm)",
         "Hệ thống tự động xếp lịch trình theo 1 trong 3 nhịp sinh học đặc thù (Vùng biển, Vùng di sản/đô thị, Vùng đồi núi), bắt buộc chèn thời gian nghỉ trưa tránh nắng (13:00 - 15:00) tại khách sạn/quán cafe máy lạnh và cố định giờ ăn uống hợp lý."),
        ("FR-05", "Tư vấn & Tối ưu hóa điểm lưu trú (Smart Hotel Anchoring & Advisory)",
         "Hệ thống đánh giá khoảng cách từ khách sạn tới cụm tham quan trong ngày. Nếu khoảng cách một chiều ≥ 35 km, hệ thống kích hoạt thẻ tư vấn với 2 phương án: Phương án A (gợi ý đổi khách sạn lân cận để tiết kiệm thời gian) và Phương án B (giữ nguyên khách sạn của khách và cảnh báo thời gian di chuyển)."),
        ("FR-06", "Tính toán chi phí & Phân bổ ngân sách chi tiết (Budget & Room Calculation)",
         "Hệ thống tự động ước tính chi phí chuyến đi dựa trên đơn giá thực tế từ cơ sở dữ liệu: tính số phòng nghỉ theo công thức ceil(Số người / 2), chia đều tiền nhiên liệu/thuê xe/vé đỗ cho cả nhóm và tính chi phí ăn uống, vé vào cổng từng người."),
        ("FR-07", "Định tuyến Road Trip đường dài & Ràng buộc an toàn sinh học (Fatigue Constraint)",
         "Khi khách di chuyển đường bộ cự ly dài (> 300 km), hệ thống tự động chèn trạm dừng nghỉ ngắn 30 phút sau mỗi 3.5 giờ lái xe liên tục, giới hạn tối đa 8 giờ lái xe/ngày và tự động đề xuất điểm dừng chân nghỉ đêm liên tỉnh (Overnight Staging City)."),
        ("FR-08", "Định tuyến trung chuyển đa phương thức & Sân bay cửa ngõ (Multi-Modal Transit)",
         "Đối với các chặng di chuyển liên tỉnh ≥ 300 km hoặc điểm đến không có sân bay thương mại (Sa Pa, Hội An, Ninh Bình...), hệ thống tự động quét các sân bay cửa ngõ trong bán kính 150 km và lập lộ trình kết hợp: Đón chặng đầu (First-mile) → Bay trung chuyển → Đường bộ chặng cuối (Last-mile)."),
        ("FR-09", "Tương tác bản đồ số & Hiển thị đường đi thực tế (Interactive Map & Polyline)",
         "Hệ thống hiển thị bản đồ trực quan bằng Leaflet/Mapbox, vẽ đường polyline thực tế bám sát mạng lưới đường bộ (thông qua Goong API cache) thay vì đường chim bay, hiển thị các marker kèm liên kết Google Place ID/URL dẫn đường đã được xác thực 100%."),
        ("FR-10", "Tùy biến & Chỉnh sửa lịch trình tương tác (Interactive Re-planning)",
         "Người dùng có thể kéo thả thay đổi thứ tự địa điểm, thêm/xóa điểm đến, ghim cố định một địa điểm vào ngày cụ thể (Hard Pinned Day), hoặc yêu cầu AI tái tính toán và giải thích lại toàn bộ lịch trình.")
    ]

    # Create Table for FR
    table_fr = doc.add_table(rows=1, cols=3)
    table_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_fr.autofit = False

    # Widths: Mã (1.0 in), Tên chức năng (2.0 in), Mô tả (3.5 in)
    col_widths = [Inches(1.0), Inches(2.0), Inches(3.5)]

    # Header Row
    hdr_cells = table_fr.rows[0].cells
    headers = ["Mã FR", "Tên chức năng", "Mô tả chi tiết"]
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(11)
        set_cell_background(hdr_cells[i], "1E40AF") # Blue 800
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)

    for idx, (code, name, desc) in enumerate(fr_data):
        row_cells = table_fr.add_row().cells
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"

        for i, val in enumerate([code, name, desc]):
            row_cells[i].width = col_widths[i]
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.bold = True
                run.font.color.rgb = RGBColor(30, 64, 175)
            elif i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.add_run(val)
            p.runs[0].font.size = Pt(10.5)
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=150, right=150)

    # SECTION III: NON-FUNCTIONAL REQUIREMENTS
    doc.add_page_break()

    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(6)
    h3_run = h3.add_run("III. 10 YÊU CẦU PHI CHỨC NĂNG (NON-FUNCTIONAL REQUIREMENTS - NFR)")
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(29, 78, 216)

    nfr_data = [
        ("NFR-01", "Hiệu năng & Thời gian phản hồi (Performance & Latency)",
         "Thời gian tính toán và trả về toàn bộ kế hoạch lịch trình đa ngày (gồm gom cụm, giải thuật TSP, tính toán ngân sách và vẽ lộ trình) không vượt quá 2.5 giây đối với thuật toán lõi, và không quá 5 giây khi có bước tạo lời giải thích từ LLM."),
        ("NFR-02", "Tính toàn vẹn & Xác thực dữ liệu (Data Integrity & Ground Truth)",
         "100% địa điểm, khách sạn và nhà hàng phải tồn tại thực tế trong cơ sở dữ liệu travel_db.db. Nghiêm cấm địa chỉ giả mạo; các liên kết Google Maps chỉ hiển thị khi có Place ID xác thực, nếu không có phải gán giá trị null để tránh người dùng lạc đường."),
        ("NFR-03", "Độ tin cậy & Khả năng chịu lỗi (Reliability & Resilience)",
         "Hệ thống áp dụng cơ chế phòng thủ đa tầng: nếu dịch vụ bên thứ ba (Goong API, Claude API) bị timeout hoặc cạn kiệt quota, hệ thống tự động chuyển sang giải thuật tính khoảng cách Haversine nội bộ và bộ đệm bộ nhớ, tuyệt đối không để xảy ra lỗi máy chủ nội bộ (HTTP 500)."),
        ("NFR-04", "Tính bảo mật (Security & API Protection)",
         "Các API Key trả phí (Google Maps API, Goong API, VyceAI LLM Key) bắt buộc phải lưu trữ an toàn ở biến môi trường backend, tuyệt đối không lộ ở mã nguồn frontend. Mọi yêu cầu gọi geocoding hoặc ma trận lộ trình đều phải đi qua Gateway API nội bộ."),
        ("NFR-05", "Khả năng mở rộng & Kiến trúc phân lớp (Scalability)",
         "Hệ thống tuân thủ mô hình AI Sandwich và Clean Architecture, phân tách độc lập giữa tầng NLU, tầng tính toán toán học cốt lõi (Deterministic Core Engine) và tầng giải thích NLG. Cho phép mở rộng dữ liệu từ 34 tỉnh hiện tại lên 63 tỉnh/thành phố mà không cần sửa cấu trúc mã nguồn."),
        ("NFR-06", "Tính chính xác trong tính toán (Mathematical Rigor)",
         "Toàn bộ bài toán định tuyến cự ly, thời gian di chuyển, số lượng phòng khách sạn và chia tiền nhóm phải được tính bằng giải thuật toán học tất định (Deterministic Python Engine), tuyệt đối không để LLM tự suy diễn hoặc nhẩm số học."),
        ("NFR-07", "Khả năng tương thích & Đa nền tảng (Compatibility)",
         "Giao diện người dùng (Frontend) tương thích hoàn toàn trên các trình duyệt hiện đại (Chrome, Edge, Firefox, Safari) và hiển thị chuẩn Responsive trên cả thiết bị di động (Mobile) lẫn máy tính để bàn (Desktop)."),
        ("NFR-08", "Tính khả dụng & Trải nghiệm người dùng (Usability)",
         "Giao diện trình bày trực quan, thông tin chi phí và lịch trình được trực quan hóa bằng dòng thời gian (Timeline) và bản đồ số đồng bộ; người dùng có thể tạo một lịch trình hoàn chỉnh chỉ với tối đa 3 thao tác cơ bản."),
        ("NFR-09", "Khả năng lưu trữ đệm & Tối ưu băng thông (Caching)",
         "Hệ thống triển khai bộ nhớ đệm kép (In-memory LRU Cache kết hợp đệm tệp MD5) cho các truy vấn lộ trình và Geocoding nhằm giảm ít nhất 70% số lượng request lặp lại ra dịch vụ ngoài, giúp tiết kiệm chi phí và tăng tốc độ phản hồi."),
        ("NFR-10", "Khả năng bảo trì & Giám sát (Maintainability & Observability)",
         "Mã nguồn được chuẩn hóa theo tiêu chuẩn PEP8, có tài liệu API tự động (OpenAPI/Swagger tại /docs), cấu trúc module hóa cao và ghi nhật ký hệ thống (structured logging) theo từng chặng tính toán để dễ dàng truy vết và gỡ lỗi.")
    ]

    # Create Table for NFR
    table_nfr = doc.add_table(rows=1, cols=3)
    table_nfr.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_nfr.autofit = False

    # Header Row
    hdr_cells_nfr = table_nfr.rows[0].cells
    headers_nfr = ["Mã NFR", "Tiêu chí chất lượng", "Mô tả chi tiết"]
    for i, title in enumerate(headers_nfr):
        hdr_cells_nfr[i].width = col_widths[i]
        p = hdr_cells_nfr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(11)
        set_cell_background(hdr_cells_nfr[i], "047857") # Emerald 700
        set_cell_margins(hdr_cells_nfr[i], top=120, bottom=120, left=150, right=150)

    for idx, (code, name, desc) in enumerate(nfr_data):
        row_cells = table_nfr.add_row().cells
        bg_color = "F0FDF4" if idx % 2 == 1 else "FFFFFF" # Light emerald tint

        for i, val in enumerate([code, name, desc]):
            row_cells[i].width = col_widths[i]
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(val)
                run.bold = True
                run.font.color.rgb = RGBColor(4, 120, 87)
            elif i == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.add_run(val)
            p.runs[0].font.size = Pt(10.5)
            set_cell_background(row_cells[i], bg_color)
            set_cell_margins(row_cells[i], top=100, bottom=100, left=150, right=150)

    output_path = r"d:\vn-travel-planner\docs\requirements_specification.docx"
    doc.save(output_path)
    print(f"Document saved successfully to {output_path}")

if __name__ == "__main__":
    create_document()
